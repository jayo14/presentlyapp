import csv
import io
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from institutions.models import Course, Attendance, AttendanceSession, CourseEnrollment

def generate_attendance_csv(course_id, date_from=None, date_to=None):
    course = Course.objects.get(id=course_id)
    sessions = AttendanceSession.objects.filter(course=course)
    if date_from:
        sessions = sessions.filter(start_time__date__gte=date_from)
    if date_to:
        sessions = sessions.filter(start_time__date__lte=date_to)
    
    sessions = sessions.order_by('start_time')
    enrollments = CourseEnrollment.objects.filter(course=course, is_active=True).select_related('student')

    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    header = ['Student Name', 'Matric Number']
    for session in sessions:
        header.append(session.start_time.strftime('%Y-%m-%d'))
    header.append('Total Present')
    header.append('Percentage')
    writer.writerow(header)

    for enrollment in enrollments:
        student = enrollment.student
        row = [student.full_name, student.matric_number]
        present_count = 0
        for session in sessions:
            is_present = Attendance.objects.filter(session=session, student=student, status='PRESENT').exists()
            row.append('P' if is_present else 'A')
            if is_present:
                present_count += 1
        
        row.append(present_count)
        percentage = (present_count / len(sessions) * 100) if sessions else 0
        row.append(f"{percentage:.1f}%")
        writer.writerow(row)
    
    return output.getvalue()

def generate_attendance_excel(course_id, date_from=None, date_to=None):
    course = Course.objects.get(id=course_id)
    sessions = AttendanceSession.objects.filter(course=course)
    if date_from:
        sessions = sessions.filter(start_time__date__gte=date_from)
    if date_to:
        sessions = sessions.filter(start_time__date__lte=date_to)
    
    sessions = sessions.order_by('start_time')
    enrollments = CourseEnrollment.objects.filter(course=course, is_active=True).select_related('student')

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"
    
    # Title row
    ws.merge_cells('A1:E1')
    ws['A1'] = f"Attendance Register - {course.code}: {course.title}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    # Header Row
    headers = ['Student Name', 'Matric Number']
    for session in sessions:
        headers.append(session.start_time.strftime('%b %d'))
    headers.extend(['Total', '%'])
    
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="F5C518", end_color="F5C518", fill_type="solid")
    
    # Data Rows
    for enrollment in enrollments:
        student = enrollment.student
        present_count = 0
        row_data = [student.full_name, student.matric_number]
        
        for session in sessions:
            is_present = Attendance.objects.filter(session=session, student=student, status='PRESENT').exists()
            row_data.append('P' if is_present else 'A')
            if is_present:
                present_count += 1
        
        row_data.append(present_count)
        percentage = (present_count / len(sessions)) if sessions else 0
        row_data.append(percentage)
        
        ws.append(row_data)
        
        # Color coding for percentage
        percent_cell = ws.cell(row=ws.max_row, column=ws.max_column)
        percent_cell.number_format = '0.0%'
        if percentage < 0.75:
            percent_cell.font = Font(color="EF4444")
        else:
            percent_cell.font = Font(color="22C55E")

    # Column widths
    for i, column_cells in enumerate(ws.columns):
        ws.column_dimensions[column_cells[0].column_letter].width = 15

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def generate_attendance_pdf(course_id, date_from=None, date_to=None):
    course = Course.objects.get(id=course_id)
    sessions = AttendanceSession.objects.filter(course=course)
    if date_from:
        sessions = sessions.filter(start_time__date__gte=date_from)
    if date_to:
        sessions = sessions.filter(start_time__date__lte=date_to)
    
    sessions = sessions.order_by('start_time')
    enrollments = CourseEnrollment.objects.filter(course=course, is_active=True).select_related('student')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], alignment=1, textColor=colors.HexColor("#F5C518"))
    elements.append(Paragraph(f"PRESENTLY ATTENDANCE REPORT", title_style))
    elements.append(Spacer(1, 12))
    
    # Course Info
    info = [
        [f"Course: {course.title} ({course.code})", f"Lecturer: {course.lecturer.full_name if course.lecturer else 'N/A'}"],
        [f"Date Range: {date_from if date_from else 'All'} to {date_to if date_to else 'All'}", f"Total Sessions: {len(sessions)}"]
    ]
    info_table = Table(info, colWidths=[250, 250])
    info_table.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 20))

    # Attendance Table
    data = [['Student Name', 'Matric Number', 'Attended', '%']]
    for enrollment in enrollments:
        student = enrollment.student
        count = Attendance.objects.filter(session__in=sessions, student=student, status='PRESENT').count()
        perc = (count / len(sessions) * 100) if sessions else 0
        data.append([student.full_name, student.matric_number, str(count), f"{perc:.1f}%"])

    t = Table(data, colWidths=[180, 120, 80, 80])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F5C518")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey)
    ]))
    elements.append(t)
    
    # Footer
    elements.append(Spacer(1, 40))
    elements.append(Paragraph("__________________________", styles['Normal']))
    elements.append(Paragraph("Lecturer Signature", styles['Normal']))

    doc.build(elements)
    return buffer.getvalue()
